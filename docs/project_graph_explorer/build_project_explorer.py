"""Project graph explorer: interactive HTML overview of the NSBI Dash web app code.

Static analysis only (Python `ast`): the app is NOT imported or run.
Extracts
  - folder / file structure (line counts, docstrings)
  - functions and their signatures, docstrings, calls (who calls whom)
  - imports between modules and to external libraries
  - Dash pages (register_page) and callbacks: Input / Output / State component.property
  - Dash component ids defined in layouts (incl. via helper functions)
  - data files read (pd.read_excel)
and injects them into explorer_template.html -> NSBI_project_explorer.html

Location in the repository:  docs/project_graph_explorer/
The project root is found automatically (two folders up). Python code under docs/
(documentation tooling, e.g. this script) is analysed too but tagged as "tooling": the
explorer hides it by default and shows it with the "Include docs/ tooling code" checkbox.

Usage (from the project root, in the PyCharm Terminal):
    .venv\Scripts\python.exe docs/project_graph_explorer/build_project_explorer.py
Options:
    --root DIR          project root (default: two folders above this script)
    --out FILE          output HTML (default: NSBI_project_explorer.html next to this script)
    --include-docs      show the docs/ tooling code in the graphs by default
"""
from __future__ import annotations

import argparse
import ast
import datetime as dt
import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
DEFAULT_ROOT = HERE.parents[1]  # docs/project_graph_explorer -> project root
TOOLING_PREFIXES = ("docs/",)  # code that is documentation tooling, not part of the app
IGNORE_DIRS = {".git", ".venv", "venv", "env", ".idea", ".vscode", "__pycache__", ".pytest_cache"}
TEXT_EXT = {".py", ".md", ".txt", ".css", ".js", ".html", ".toml", ".cfg", ".ini", ".yml", ".yaml", ".json", ""}
CALLBACK_DECORATORS = {"callback", "dash.callback", "app.callback", "clientside_callback"}
DEP_KINDS = {"Input": "input", "Output": "output", "State": "state"}


def first_line(doc: str | None) -> str:
    return (doc or "").strip().split("\n\n")[0].replace("\n", " ").strip()


def unparse(node) -> str:
    try:
        return ast.unparse(node)
    except Exception:  # pragma: no cover
        return "?"


def module_name(rel: Path) -> str:
    parts = list(rel.with_suffix("").parts)
    if parts[-1] == "__init__":
        parts = parts[:-1]
    return ".".join(parts)


# ----------------------------------------------------------------------------- file tree
def scan_tree(root: Path) -> tuple[dict, list[dict]]:
    files = []

    def walk(p: Path) -> dict:
        node = {"name": p.name, "path": p.relative_to(root).as_posix() if p != root else "", "type": "dir", "children": []}
        for c in sorted(p.iterdir(), key=lambda x: (x.is_file(), x.name.lower())):
            if c.is_dir():
                if c.name in IGNORE_DIRS:
                    continue
                node["children"].append(walk(c))
            else:
                rel = c.relative_to(root).as_posix()
                info = {"name": c.name, "path": rel, "type": "file", "ext": c.suffix.lower(), "size": c.stat().st_size}
                if c.suffix.lower() in TEXT_EXT:
                    try:
                        info["lines"] = len(c.read_text(encoding="utf-8").splitlines())
                    except UnicodeDecodeError:
                        pass
                node["children"].append(info)
                files.append(info)
        return node

    return walk(root), files


# ----------------------------------------------------------------------------- python analysis
class FileAnalyser:
    def __init__(self, root: Path, path: Path, project_modules: set[str]):
        self.root, self.path = root, path
        self.rel = path.relative_to(root).as_posix()
        self.module = module_name(path.relative_to(root))
        self.project_modules = project_modules
        self.src = path.read_text(encoding="utf-8")
        self.tree = ast.parse(self.src)
        self.aliases: dict[str, tuple[str, str | None, bool]] = {}  # local -> (module, name, is_project)
        self.imports: list[dict] = []
        self.functions: dict[str, dict] = {}
        self.page: dict | None = None
        self.component_defs: list[dict] = []
        self.constants: list[str] = []

    def is_project(self, mod: str) -> bool:
        return mod in self.project_modules or any(mod.startswith(m + ".") for m in self.project_modules)

    def run(self):
        self._imports()
        self._defs(self.tree.body, prefix="")
        self._module_level()
        return self

    def _imports(self):
        for n in ast.walk(self.tree):
            if isinstance(n, ast.Import):
                for a in n.names:
                    proj = self.is_project(a.name)
                    self.aliases[a.asname or a.name.split(".")[0]] = (a.name if a.asname else a.name.split(".")[0], None, proj)
                    self.imports.append({"module": a.name, "names": [], "project": proj, "line": n.lineno})
            elif isinstance(n, ast.ImportFrom) and n.module:
                proj = self.is_project(n.module)
                names = []
                for a in n.names:
                    sub = f"{n.module}.{a.name}"
                    if sub in self.project_modules:  # from package import module
                        self.aliases[a.asname or a.name] = (sub, None, True)
                    else:
                        self.aliases[a.asname or a.name] = (n.module, a.name, proj)
                    names.append(a.name)
                if n.module != "__future__":
                    self.imports.append({"module": n.module, "names": names, "project": proj, "line": n.lineno})

    def _defs(self, body, prefix: str, parent: str | None = None):
        for n in body:
            if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)):
                q = f"{prefix}{n.name}"
                self.functions[q] = self._function_info(n, q, parent)
                self._defs(n.body, prefix=q + ".", parent=q)
            elif isinstance(n, ast.ClassDef):
                q = f"{prefix}{n.name}"
                self.functions[q] = {"name": n.name, "qualname": q, "kind": "class", "line": n.lineno,
                                     "end": n.end_lineno, "doc": first_line(ast.get_docstring(n)),
                                     "signature": f"class {n.name}", "decorators": [unparse(d) for d in n.decorator_list],
                                     "parent": parent, "callback": None, "_node": n}
            elif not prefix and isinstance(n, (ast.Assign, ast.AnnAssign)):
                targets = n.targets if isinstance(n, ast.Assign) else [n.target]
                for t in targets:
                    if isinstance(t, ast.Name) and t.id.isupper():
                        self.constants.append(t.id)

    def _function_info(self, n, q, parent):
        args = unparse(n.args)
        ret = f" -> {unparse(n.returns)}" if n.returns else ""
        info = {"name": n.name, "qualname": q, "kind": "function", "line": n.lineno, "end": n.end_lineno,
                "doc": first_line(ast.get_docstring(n)), "signature": f"def {n.name}({args}){ret}",
                "params": [a.arg for a in n.args.args + n.args.kwonlyargs],
                "decorators": [unparse(d) for d in n.decorator_list], "parent": parent, "callback": None, "_node": n}
        for d in n.decorator_list:
            if isinstance(d, ast.Call) and unparse(d.func) in CALLBACK_DECORATORS:
                deps, opts = [], {}
                for a in d.args:
                    if isinstance(a, ast.Call) and unparse(a.func).split(".")[-1] in DEP_KINDS:
                        cid = a.args[0].value if a.args and isinstance(a.args[0], ast.Constant) else unparse(a.args[0]) if a.args else "?"
                        prop = a.args[1].value if len(a.args) > 1 and isinstance(a.args[1], ast.Constant) else "?"
                        dup = any(k.arg == "allow_duplicate" for k in a.keywords)
                        deps.append({"kind": DEP_KINDS[unparse(a.func).split(".")[-1]], "id": cid, "prop": prop,
                                     "allow_duplicate": dup})
                for k in d.keywords:
                    opts[k.arg] = unparse(k.value)
                info["kind"] = "callback"
                info["callback"] = {"deps": deps, "options": opts}
        return info

    # --- calls -------------------------------------------------------------------------
    def resolve_call(self, func, scope: str | None) -> dict | None:
        """Return {'type': 'project'|'external', ...} for a call target, or None."""
        if isinstance(func, ast.Name):
            name = func.id
            # nested function visible from this scope (or its parents)
            s = scope
            while s:
                if f"{s}.{name}" in self.functions:
                    return {"type": "project", "module": self.module, "qualname": f"{s}.{name}"}
                s = self.functions.get(s, {}).get("parent")
            if name in self.functions:
                return {"type": "project", "module": self.module, "qualname": name}
            if name in self.aliases:
                mod, obj, proj = self.aliases[name]
                if proj:
                    return {"type": "project", "module": mod, "qualname": obj or name}
                return {"type": "external", "lib": mod.split(".")[0], "name": f"{mod}.{obj}" if obj else mod}
            return None
        if isinstance(func, ast.Attribute):
            chain = []
            cur = func
            while isinstance(cur, ast.Attribute):
                chain.append(cur.attr)
                cur = cur.value
            if isinstance(cur, ast.Name) and cur.id in self.aliases:
                mod, obj, proj = self.aliases[cur.id]
                full = ".".join([mod] + ([obj] if obj else []) + chain[::-1])
                if proj:
                    return {"type": "project", "module": mod if not obj else mod, "qualname": chain[0] if not obj else obj}
                return {"type": "external", "lib": mod.split(".")[0], "name": full}
        return None

    def calls_in(self, node, scope):
        found = []

        def visit(n):
            for child in ast.iter_child_nodes(n):
                if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    continue  # nested defs analysed separately
                if isinstance(child, ast.Call):
                    r = self.resolve_call(child.func, scope)
                    if r:
                        r["line"] = child.lineno
                        found.append(r)
                visit(child)

        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            for b in node.body:
                if isinstance(b, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
                    continue
                wrapper = ast.Module(body=[b], type_ignores=[])
                visit(wrapper)
        else:
            visit(node)
        return found

    def _module_level(self):
        body = [b for b in self.tree.body if not isinstance(b, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef,
                                                                  ast.Import, ast.ImportFrom))]
        self.module_calls = self.calls_in(ast.Module(body=body, type_ignores=[]), None)
        for n in ast.walk(self.tree):
            if isinstance(n, ast.Call) and unparse(n.func) in ("dash.register_page", "register_page"):
                self.page = {k.arg: (k.value.value if isinstance(k.value, ast.Constant) else unparse(k.value))
                             for k in n.keywords}


def find_component_defs(an: FileAnalyser, known_ids: set[str], helpers: dict) -> list[dict]:
    """Component ids created in this file: id="x" keyword, or helper(...) calls passing the id."""
    defs = []

    def enclosing(lineno):
        best = None
        for q, f in an.functions.items():
            if f["line"] <= lineno <= (f["end"] or f["line"]):
                if best is None or f["line"] >= an.functions[best]["line"]:
                    best = q
        return best

    for n in ast.walk(an.tree):
        if not isinstance(n, ast.Call):
            continue
        fname = unparse(n.func)
        for k in n.keywords:
            if k.arg == "id" and isinstance(k.value, ast.Constant) and isinstance(k.value.value, str):
                defs.append({"id": k.value.value, "type": fname, "line": n.lineno, "func": enclosing(n.lineno)})
        target = an.resolve_call(n.func, None)
        if target and target["type"] == "project":
            key = f"{target['module']}::{target['qualname']}"
            if key in helpers:
                pidx, ctype = helpers[key]
                if pidx < len(n.args) and isinstance(n.args[pidx], ast.Constant):
                    defs.append({"id": n.args[pidx].value, "type": f"{ctype} (via {target['qualname']}())",
                                 "line": n.lineno, "func": enclosing(n.lineno)})
    return defs


def find_helpers(an: FileAnalyser) -> dict:
    """Project functions that create a component whose id is one of their parameters."""
    out = {}
    for q, f in an.functions.items():
        if f["kind"] == "class":
            continue
        for n in ast.walk(f["_node"]):
            if isinstance(n, ast.Call):
                for k in n.keywords:
                    if k.arg == "id" and isinstance(k.value, ast.Name) and k.value.id in f["params"]:
                        out[f"{an.module}::{q}"] = (f["params"].index(k.value.id), unparse(n.func))
    return out


def excel_sheets(path: Path) -> list[dict]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True)
        res = []
        for ws in wb.worksheets:
            header = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            res.append({"name": ws.title, "rows": max(ws.max_row - 1, 0), "cols": ws.max_column,
                        "header": [h for h in header if h is not None][:80]})
        return res
    except Exception as e:  # openpyxl missing or unreadable
        return [{"name": f"(could not read: {e})", "rows": 0, "cols": 0, "header": []}]


def is_tooling(rel: str) -> bool:
    return rel.startswith(TOOLING_PREFIXES)


def analyse(root: Path, include_docs: bool = False) -> dict:
    tree, files = scan_tree(root)
    py_files = [root / f["path"] for f in files if f["ext"] == ".py"]
    project_modules = {module_name(p.relative_to(root)) for p in py_files}
    project_modules |= {m.rsplit(".", 1)[0] for m in project_modules if "." in m}
    analysers = {}
    for p in py_files:
        a = FileAnalyser(root, p, project_modules).run()
        analysers[a.module] = a
    mod2file = {m: a.rel for m, a in analysers.items()}

    helpers = {}
    for a in analysers.values():
        helpers.update(find_helpers(a))

    functions, calls, callbacks, ext_libs = [], [], [], {}

    def fid(module, qual):
        return f"fn:{mod2file.get(module, module)}::{qual}"

    def add_calls(src_id, found, file_rel):
        for c in found:
            if c["type"] == "project":
                tgt_mod = c["module"]
                if tgt_mod not in analysers:
                    continue
                if c["qualname"] not in analysers[tgt_mod].functions:
                    continue
                calls.append({"source": src_id, "target": fid(tgt_mod, c["qualname"]), "line": c["line"], "kind": "call"})
            else:
                lib = c["lib"]
                ext_libs.setdefault(lib, {"name": lib, "used_by": set(), "calls": set()})
                ext_libs[lib]["used_by"].add(file_rel)
                ext_libs[lib]["calls"].add(c["name"])
                calls.append({"source": src_id, "target": f"ext:{lib}", "line": c["line"], "kind": "ext", "name": c["name"]})

    data_files = [f for f in files if f["ext"] in (".xlsx", ".xls", ".csv", ".parquet", ".json") and f["path"].startswith("data/")]
    data_reads = []
    for m, a in analysers.items():
        for q, f in a.functions.items():
            node_id = fid(m, q)
            functions.append({k: v for k, v in f.items() if k != "_node"}
                             | {"id": node_id, "file": a.rel, "module": m, "tooling": is_tooling(a.rel)})
            found = a.calls_in(f["_node"], q)
            add_calls(node_id, found, a.rel)
            if any(c["type"] == "external" and c["name"].split(".")[-1] in ("read_excel", "read_csv", "read_parquet")
                   for c in found):
                for df in data_files:
                    data_reads.append({"source": f"data:{df['path']}", "target": node_id, "kind": "read"})
            if f["callback"]:
                callbacks.append({"id": node_id, "file": a.rel, "name": q, **f["callback"]})
        if a.module_calls:
            node_id = fid(m, "<module>")
            functions.append({"id": node_id, "file": a.rel, "module": m, "tooling": is_tooling(a.rel),
                              "name": "<module level>", "qualname": "<module>",
                              "kind": "module", "line": 1, "end": None, "doc": "Code executed when the file is imported",
                              "signature": "", "decorators": [], "parent": None, "callback": None})
            add_calls(node_id, a.module_calls, a.rel)

    # module import edges
    imports = []
    for m, a in analysers.items():
        for imp in a.imports:
            if imp["project"]:
                tmod = imp["module"]
                targets = []
                for nm in imp["names"] or [None]:
                    sub = f"{tmod}.{nm}" if nm else None
                    if sub and sub in mod2file:
                        targets.append((sub, [nm]))
                    elif tmod in mod2file:
                        targets.append((tmod, [nm] if nm else []))
                for t, names in targets:
                    imports.append({"source": f"file:{a.rel}", "target": f"file:{mod2file[t]}", "names": names, "line": imp["line"]})
            else:
                lib = imp["module"].split(".")[0]
                ext_libs.setdefault(lib, {"name": lib, "used_by": set(), "calls": set()})
                ext_libs[lib]["used_by"].add(a.rel)
                imports.append({"source": f"file:{a.rel}", "target": f"ext:{lib}", "names": imp["names"], "line": imp["line"]})
    # merge duplicate import edges
    merged = {}
    for e in imports:
        k = (e["source"], e["target"])
        if k in merged:
            merged[k]["names"] = sorted(set(merged[k]["names"]) | set(e["names"]))
        else:
            merged[k] = dict(e)
    imports = list(merged.values())

    # framework conventions (not visible as imports)
    framework = []
    app_file = next((a.rel for a in analysers.values() if not is_tooling(a.rel) and "use_pages=True" in a.src), None)
    if app_file:
        for a in analysers.values():
            if a.page:
                framework.append({"source": f"file:{app_file}", "target": f"file:{a.rel}", "kind": "pages",
                                  "label": "use_pages: page registered"})
        for f in files:
            if f["path"].startswith("assets/"):
                framework.append({"source": f"file:{app_file}", "target": f"file:{f['path']}", "kind": "assets",
                                  "label": "auto-loaded from assets/"})

    # Dash components
    known_ids = {d["id"] for cb in callbacks for d in cb["deps"]}
    components = {}
    for a in analysers.values():
        if is_tooling(a.rel):
            continue
        for d in find_component_defs(a, known_ids, helpers):
            c = components.setdefault(d["id"], {"id": d["id"], "defs": [], "inputs": [], "outputs": [], "states": []})
            c["defs"].append({"file": a.rel, "type": d["type"], "line": d["line"], "func": d["func"]})
    for cb in callbacks:
        for d in cb["deps"]:
            c = components.setdefault(d["id"], {"id": d["id"], "defs": [], "inputs": [], "outputs": [], "states": []})
            c[d["kind"] + "s"].append({"callback": cb["id"], "prop": d["prop"]})

    # data files + (optional) external scripts writing them
    data_nodes = []
    for df in data_files:
        p = root / df["path"]
        data_nodes.append({"id": f"data:{df['path']}", "path": df["path"], "size": df["size"],
                           "sheets": excel_sheets(p) if df["ext"] in (".xlsx", ".xls") else []})
    writers = []  # kept for the template (scripts writing data files); none at the moment

    file_info = {}
    for m, a in analysers.items():
        file_info[a.rel] = {"module": m, "doc": first_line(ast.get_docstring(a.tree)), "page": a.page,
                            "imports": a.imports, "constants": a.constants,
                            "functions": [q for q in a.functions], "tooling": is_tooling(a.rel)}

    return {
        "generated": dt.datetime.now().strftime("%Y-%m-%d %H:%M"),
        "root": root.name,
        "include_docs_default": include_docs,
        "tooling_prefixes": list(TOOLING_PREFIXES),
        "tree": tree,
        "files": files,
        "file_info": file_info,
        "functions": functions,
        "calls": calls,
        "imports": imports,
        "framework": framework,
        "callbacks": callbacks,
        "components": list(components.values()),
        "external": [{"name": v["name"], "used_by": sorted(v["used_by"]), "calls": sorted(v["calls"])}
                     for v in sorted(ext_libs.values(), key=lambda x: x["name"])],
        "data_files": data_nodes,
        "data_reads": data_reads,
        "writers": writers,
    }


def main():
    ap = argparse.ArgumentParser(description="Build the NSBI project graph explorer (HTML).")
    ap.add_argument("--root", type=Path, default=DEFAULT_ROOT, help="project root folder")
    ap.add_argument("--out", type=Path, default=HERE / "NSBI_project_explorer.html", help="output HTML file")
    ap.add_argument("--include-docs", action="store_true",
                    help="show the docs/ tooling code in the graphs by default (checkbox pre-ticked)")
    args = ap.parse_args()
    root = args.root.resolve()
    if not (root / "app.py").exists():
        sys.exit(f"app.py not found in {root}. Use --root to point at the project folder.")
    data = analyse(root, include_docs=args.include_docs)
    template = (HERE / "explorer_template.html").read_text(encoding="utf-8")
    payload = json.dumps(data, ensure_ascii=False, default=list).replace("</", "<\\/")
    args.out.write_text(template.replace("/*__DATA__*/null", payload), encoding="utf-8")
    print(f"Wrote {args.out}  ({len(data['functions'])} functions, {len(data['callbacks'])} callbacks, "
          f"{len(data['components'])} components, {len(data['files'])} files)")


if __name__ == "__main__":
    main()
